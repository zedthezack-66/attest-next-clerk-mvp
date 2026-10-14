import io
from dataclasses import dataclass
from pathlib import Path

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".png", ".jpg", ".jpeg", ".tif", ".tiff", ".webp"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".webp"}


@dataclass
class ExtractedText:
    text: str
    # direct_text | ocr | ocr_unavailable | unsupported
    method: str
    confidence: float


def _ocr_image(data: bytes) -> str | None:
    """Tesseract OCR when installed (pip install pytesseract pillow + tesseract binary)."""
    try:
        import pytesseract
        from PIL import Image
    except ImportError:
        return None
    try:
        return pytesseract.image_to_string(Image.open(io.BytesIO(data)))
    except Exception:  # noqa: BLE001 - missing tesseract binary surfaces as a generic error
        return None


def _confidence(text: str, base: float) -> float:
    if len(text.strip()) < 80:
        base -= 0.25
    return round(max(0.1, base), 2)


def extract_text(filename: str, data: bytes) -> ExtractedText:
    suffix = Path(filename).suffix.lower()
    if suffix == ".txt":
        text = data.decode("utf-8", errors="replace")
        return ExtractedText(text, "direct_text", _confidence(text, 0.98))
    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
        if text.strip():
            return ExtractedText(text, "direct_text", _confidence(text, 0.93))
        return ExtractedText("", "ocr_unavailable", 0.0)  # scanned PDF: needs a rasterise + OCR adapter
    if suffix == ".docx":
        from docx import Document as DocxDocument

        doc = DocxDocument(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        text = "\n".join(parts)
        return ExtractedText(text, "direct_text", _confidence(text, 0.93))
    if suffix in IMAGE_EXTENSIONS:
        text = _ocr_image(data)
        if text is None:
            return ExtractedText("", "ocr_unavailable", 0.0)
        return ExtractedText(text, "ocr", _confidence(text, 0.72))
    return ExtractedText("", "unsupported", 0.0)
